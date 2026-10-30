# -*- coding: utf-8 -*-
"""
Prueba de humo de los reportes.

Genera la base wide y los seis reportes (Excel, Word y PowerPoint de ingreso y
de seguimiento) con una base sintética de cada país, por el mismo camino que
usa la app: renombrar con RENAME_MAP, `procesar_wide()` y `run_script()`.

No revisa las cifras, eso lo hacen las pruebas de decisiones. Revisa que cada
reporte se arme, que el archivo se pueda abrir y que traiga contenido. Existe
porque los cuatro defectos del 2026-09-09 los encontró Ecuador y no las
pruebas: ninguna generaba un reporte.

Tarda alrededor de un minuto, porque el Word y el PowerPoint corren en un
proceso aparte, igual que en la app.
"""
import io
import os
import sys
import tempfile

import pandas as pd
import pytest

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from datos_sinteticos import base_sintetica, excel_bruto  # noqa: E402
from pipeline.runner import OUTPUTS, es_falta_de_datos, run_script  # noqa: E402
from pipeline.validacion_top import CATEGORIAS_POR_PAIS  # noqa: E402
from pipeline.wide_top import procesar_wide  # noqa: E402

PAISES = ['Perú', 'Ecuador', 'El Salvador', 'México']
REPORTES = list(OUTPUTS)


def _wide(df, **filtros):
    """Base bruta → base wide en disco, como hace la app."""
    carpeta = tempfile.mkdtemp(prefix='qalat_humo_')
    raw = excel_bruto(df, os.path.join(carpeta, 'bruto.xlsx'))
    res = procesar_wide(raw, **filtros)
    ruta = os.path.join(carpeta, 'TOP_Base_Wide.xlsx')
    with open(ruta, 'wb') as fh:
        fh.write(res['excel_bytes'].getvalue())
    return ruta


def _texto(datos, nombre):
    """Todo el texto de un .xlsx, .docx o .pptx, para comprobar que trae contenido."""
    ext = nombre.rsplit('.', 1)[1]
    if ext == 'xlsx':
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(datos), read_only=True)
        return ' '.join(str(v) for ws in wb.worksheets
                        for fila in ws.iter_rows(values_only=True) for v in fila if v is not None)
    if ext == 'docx':
        from docx import Document
        doc = Document(io.BytesIO(datos))
        celdas = [c.text for t in doc.tables for f in t.rows for c in f.cells]
        return ' '.join([p.text for p in doc.paragraphs] + celdas)
    from pptx import Presentation
    prs = Presentation(io.BytesIO(datos))
    assert len(prs.slides) >= 3, f'{nombre}: solo {len(prs.slides)} diapositivas'
    return ' '.join(sh.text_frame.text for s in prs.slides for sh in s.shapes if sh.has_text_frame)


@pytest.fixture(scope='module', params=PAISES)
def wide_pais(request):
    return request.param, _wide(base_sintetica(request.param))


@pytest.mark.parametrize('reporte', REPORTES)
def test_cada_reporte_se_arma_en_cada_pais(wide_pais, reporte):
    pais, ruta = wide_pais
    buf, nombre, _ = run_script(reporte, ruta)
    datos = buf.getvalue()
    assert len(datos) > 5000, f'{pais} · {nombre} salió casi vacío'
    texto = _texto(datos, nombre)
    if reporte != 'pptx_seg':   # el PPT de seguimiento usa gráficos, no texto
        sustancias = [c for c in CATEGORIAS_POR_PAIS[pais] if c in texto]
        assert sustancias, f'{pais} · {nombre} no menciona ninguna sustancia del país'


@pytest.mark.parametrize('reporte', REPORTES)
def test_los_reportes_de_un_centro(reporte):
    """El camino de la pestaña Reportes del centro: wide filtrada y filtro_centro."""
    ruta = _wide(base_sintetica('Ecuador'), filtro_centro='C01')
    buf, nombre, _ = run_script(reporte, ruta, filtro_centro='C01')
    assert nombre.endswith(f'_C01.{nombre.rsplit(".", 1)[1]}')
    assert len(buf.getvalue()) > 5000


def test_el_filtro_de_periodo_no_rompe_la_wide():
    """Defecto 3 del 2026-09-09: KeyError '_episodio' con filtro de periodo."""
    ruta = _wide(base_sintetica('Ecuador'), fecha_desde='2026-03', fecha_hasta='2026-07')
    buf, _, _ = run_script('caract_excel', ruta)
    assert len(buf.getvalue()) > 5000


@pytest.mark.parametrize('reporte', ['pdf_caract', 'pptx_caract', 'caract_excel'])
def test_un_centro_donde_nadie_consume_otra_sustancia(reporte):
    """Defecto 1 del 2026-09-09: una sustancia sin consumidores llegaba con prom=None."""
    ruta = _wide(base_sintetica('Ecuador', n_pacientes=6, centros=('C01',), solo_alcohol=True))
    buf, _, _ = run_script(reporte, ruta)
    assert len(buf.getvalue()) > 5000


@pytest.mark.parametrize('reporte', REPORTES)
def test_un_centro_sin_top_de_ingreso_avisa_y_no_revienta(reporte):
    """Defecto 2 del 2026-09-09. Sin ingreso no hay reporte, pero el motivo tiene
    que llegar como falta de datos, no como error técnico."""
    df = base_sintetica('Ecuador', n_pacientes=4, centros=('C01',), etapa_ingreso='en_tratamiento')
    ruta = _wide(df)
    with pytest.raises(Exception) as info:
        run_script(reporte, ruta)
    assert es_falta_de_datos(info.value), f'error técnico en vez de aviso: {str(info.value)[-300:]}'


@pytest.mark.parametrize('reporte', ['seg_excel', 'pdf_seg', 'pptx_seg'])
def test_sin_ningun_seguimiento_los_reportes_de_seguimiento_avisan(reporte):
    ruta = _wide(base_sintetica('Perú', n_pacientes=8, con_seguimiento=False))
    with pytest.raises(Exception) as info:
        run_script(reporte, ruta)
    assert es_falta_de_datos(info.value), f'error técnico en vez de aviso: {str(info.value)[-300:]}'


# ── Hoja "TOP sin ingreso" de la base wide ────────────────────────────────────
# DECISIONES.md, "Los TOP sin ingreso se listan aparte". El caso que la motivó:
# un TOP de egreso de CIJ que figuraba en la revisión enviada al país y que el
# país no encontraba en su descarga, porque no tenía un TOP de ingreso antes.

def _base_con_egreso_sin_ingreso():
    base = base_sintetica('México', n_pacientes=6)
    extra = base.iloc[[0]].copy()
    extra['id'] = 999
    extra['codigo_paciente'] = 'ZZZZ01011990'
    extra['fecha_nacimiento'] = '1990-01-01'
    extra['fecha_entrevista'] = '2026-10-16'
    extra['etapa'] = 'egreso'
    return base, pd.concat([base, extra], ignore_index=True)


def _procesar(df):
    carpeta = tempfile.mkdtemp(prefix='qalat_sin_ingreso_')
    return procesar_wide(excel_bruto(df, os.path.join(carpeta, 'bruto.xlsx')))


def _hoja(res, nombre):
    from openpyxl import load_workbook
    return load_workbook(io.BytesIO(res['excel_bytes'].getvalue()))[nombre]


def test_la_hoja_top_sin_ingreso_lista_el_top_que_queda_fuera():
    base, con_egreso = _base_con_egreso_sin_ingreso()
    centro = str(base.iloc[0]['centro'])
    res = _procesar(con_egreso)
    assert [f['N.º de registro'] for f in res['sin_ingreso']] == [999]
    ws = _hoja(res, 'TOP sin ingreso')
    filas = [r for r in ws.iter_rows(min_row=4, values_only=True) if any(v is not None for v in r)]
    assert filas == [(centro, 'ZZZZ01011990', '16-10-2026', 'Egreso', 999)]
    assert '1 registro fuera' in ws['A1'].value


def test_la_hoja_es_la_septima_y_no_cambia_la_base_wide():
    base, con_egreso = _base_con_egreso_sin_ingreso()
    sin, con = _procesar(base), _procesar(con_egreso)
    from openpyxl import load_workbook
    hojas = load_workbook(io.BytesIO(con['excel_bytes'].getvalue())).sheetnames
    assert hojas == ['Base Wide', 'Resumen', 'Alertas', 'Calidad de Datos',
                     'Por Centro', 'Pendientes TOP2', 'TOP sin ingreso']
    assert len(con['wide']) == len(sin['wide'])
    assert con['stats']['N_total'] == sin['stats']['N_total']


def test_sin_casos_la_hoja_lo_dice():
    res = _procesar(base_sintetica('México', n_pacientes=6))
    assert res['sin_ingreso'] == []
    assert 'Todos los TOP tienen' in str(_hoja(res, 'TOP sin ingreso')['A4'].value)
